#!/usr/bin/env python3
"""Builds a Previo .pptx deck from a JSON outline, on the fixed grid of templates/deck.html.

    python3 scripts/build_deck.py deck.json out.pptx [--preview out.html]

Needs python-pptx (pip install python-pptx). The JSON format and slide types are described in
templates/deck-structure.md ("Generator"); scripts/deck-example.json is a complete example.

Every position comes from deck.html (1280 x 720 px, 1 px = 7143.75 EMU = 0.5625 pt), so all content
slides share one grid and nothing jumps between slides. Text is measured with the bundled Inter
metrics: content that would not fit is refused with a message saying what to shorten or split.
Fix the content, do not move the boxes. Colors follow the red rule automatically (a number slide
gets a gray label and gray icon tiles).
--preview writes an HTML rendering of the same layout, for checking without PowerPoint.
"""
import argparse, base64, json, pathlib, re, shutil, struct, subprocess, sys, tempfile, zlib

SKILL = pathlib.Path(__file__).resolve().parent.parent
EMU = 914400 / 128  # EMU per px: the 1280 px wide slide is 10 in
PT = 0.5625         # pt per px
C = dict(red="b50000", soft="faeaeb", ink="202124", muted="515151", g400="b2b2b2",
         border="e5e5e5", surface="f2f2f2", white="ffffff")
ICONS = {p.stem for p in (SKILL / "icons").glob("*.svg")}
CONTENT_TOP, CONTENT_BOTTOM = 268, 632  # content block on white slides; footer line at 656
REGULAR = (64, "red", (28, 36), (23, 32), 24, 3)  # item: tile, tile color, title and text size/line, gap, max text lines
COMPACT = (56, "red", (24, 32), (22, 30), 20, 2)  # item in two columns: smaller tile and type
DENSE = COMPACT[:-1] + (3,)  # last resort before an error: compact type, three lines of text
QUOTES = {"cs": "„“", "sk": "„“", "de": "„“", "pl": "„”"}  # other languages: “”
FOCUS = dict(center=(.5, .5), top=(.5, 0), bottom=(.5, 1), left=(0, .5), right=(1, .5))
SHARP = 1.5  # photo px per slide px: a full-slide photo needs 1920 x 1080 to stay sharp on a projector
TMP = pathlib.Path(tempfile.mkdtemp(prefix="previo-deck-"))  # scaled photos and the scrim; the preview links to them


def logo(variant):
    """Logo PNG and its width/height ratio (the variants differ slightly)."""
    path = SKILL / "logo" / "png" / f"previo-logo-{variant}.png"
    w, h = struct.unpack(">II", path.read_bytes()[16:24])
    return path, w / h


def typeset(text, lang):
    """Non-breaking spaces where a line must not break: 15 %, 2 000 Kč, and Czech one-letter words."""
    text = re.sub(r"(\d) (?=\d{3}\b|%|‰|Kč|€|\$|h\b|min\b|ks\b)", "\\1\u00a0", text)
    if lang.startswith(("cs", "sk")):
        text = re.sub(r"(?<![\w\u00a0])([kKsSvVzZoOuUaAiI]) ", "\\1\u00a0", text)
    return text


def crop(iw, ih, w, h, focus):
    """Crop fractions (left, top, right, bottom) filling a w x h frame without stretching; the focus
    point (0–1 from the left and top) stays in view. Same rule as CSS object-position."""
    fx, fy = focus
    if iw / ih > w / h:
        c = 1 - (w / h) / (iw / ih)
        return c * fx, 0, c * (1 - fx), 0
    c = 1 - (iw / ih) / (w / h)
    return 0, c * fy, 0, c * (1 - fy)


def scrim():
    """1 x 256 px PNG, clear at the top and dark at the bottom: stretched over a photo it keeps white text legible."""
    path = TMP / "scrim.png"
    if not path.exists():
        rows = b"".join(b"\0" + bytes((20, 20, 22, round(210 * y / 255))) for y in range(256))
        chunk = lambda t, d: struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d))
        path.write_bytes(b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 256, 8, 6, 0, 0, 0))
                         + chunk(b"IDAT", zlib.compress(rows)) + chunk(b"IEND", b""))
    return path


def shrink(path, w, h):
    """JPEG copy scaled to SHARP x the frame, so a deck with photos stays a few MB. Uses Pillow, or sips
    on a Mac; returns the original when neither is there or the photo is not larger than needed."""
    try:
        from PIL import Image, ImageOps
        im = ImageOps.exif_transpose(Image.open(path))  # phone photos are often stored rotated
        k = max(w * SHARP / im.width, h * SHARP / im.height)
        if k >= 1 and path.suffix.lower() in (".jpg", ".jpeg"):
            return path
        out = TMP / f"{len(list(TMP.iterdir()))}-{path.stem}.jpg"
        if k < 1:
            im = im.resize((round(im.width * k), round(im.height * k)), Image.LANCZOS)
        im.convert("RGB").save(out, quality=82, optimize=True)
        return out
    except ImportError:
        pass
    from pptx.parts.image import Image as PImage
    iw, ih = PImage.from_file(str(path)).size
    k = max(w * SHARP / iw, h * SHARP / ih)
    if k >= 1 or not shutil.which("sips"):
        return path
    out = TMP / f"{len(list(TMP.iterdir()))}-{path.stem}.jpg"
    subprocess.run(["sips", "-s", "format", "jpeg", "-s", "formatOptions", "82", "-z", str(round(ih * k)), str(round(iw * k)),
                    str(path), "--out", str(out)], check=True, capture_output=True)
    return out


def on_red(alpha):
    """White at the given opacity over primary red, as a solid color (pptx text has no opacity)."""
    return "%02x%02x%02x" % tuple(round(c + (255 - c) * alpha) for c in (181, 0, 0))


# --- text measurement (advance widths from the bundled TTFs, no extra library) -------------------

class Font:
    def __init__(self, path):
        d = path.read_bytes()
        tables = {}
        for i in range(struct.unpack(">H", d[4:6])[0]):
            tag, _, off, _ = struct.unpack(">4sIII", d[12 + 16 * i:28 + 16 * i])
            tables[tag.decode()] = off
        self.upem = struct.unpack(">H", d[tables["head"] + 18:tables["head"] + 20])[0]
        n = struct.unpack(">H", d[tables["hhea"] + 34:tables["hhea"] + 36])[0]
        h = tables["hmtx"]
        self.adv = [struct.unpack(">H", d[h + 4 * i:h + 4 * i + 2])[0] for i in range(n)]
        self.cmap = self._cmap(d, tables["cmap"])

    @staticmethod
    def _cmap(d, base):
        recs = {}
        for i in range(struct.unpack(">H", d[base + 2:base + 4])[0]):
            pid, eid, off = struct.unpack(">HHI", d[base + 4 + 8 * i:base + 12 + 8 * i])
            recs[(pid, eid)] = base + off
        out = {}
        if (3, 10) in recs:  # format 12
            o = recs[(3, 10)]
            for g in range(struct.unpack(">I", d[o + 12:o + 16])[0]):
                start, end, gid = struct.unpack(">III", d[o + 16 + 12 * g:o + 28 + 12 * g])
                if start < 0x3000:
                    for cp in range(start, min(end, 0x2fff) + 1):
                        out[cp] = gid + cp - start
            return out
        o = recs[(3, 1)]  # format 4
        seg = struct.unpack(">H", d[o + 6:o + 8])[0] // 2
        arr = lambda k: struct.unpack(f">{seg}H", d[o + k:o + k + 2 * seg])
        ends, starts = arr(14), arr(16 + 2 * seg)
        deltas, ro = arr(16 + 4 * seg), arr(16 + 6 * seg)
        for i in range(seg):
            for cp in range(starts[i], ends[i] + 1):
                if cp == 0xffff:
                    continue
                if ro[i] == 0:
                    gid = (cp + deltas[i]) & 0xffff
                else:
                    p = o + 16 + 6 * seg + 2 * i + ro[i] + 2 * (cp - starts[i])
                    gid = struct.unpack(">H", d[p:p + 2])[0]
                    gid = (gid + deltas[i]) & 0xffff if gid else 0
                out[cp] = gid
        return out

    def width(self, text, size, spacing=0):
        units = sum(self.adv[min(self.cmap.get(ord(ch), 0), len(self.adv) - 1)] for ch in text)
        return units * size / self.upem + spacing * size * len(text)


FONTS = {"reg": Font(SKILL / "fonts" / "Inter-Regular.ttf"), "med": Font(SKILL / "fonts" / "Inter-Medium.ttf")}
FAMILY = {"reg": "Inter", "med": "Inter Medium"}


def count_lines(text, size, width, font="reg", spacing=0):
    f, limit, n = FONTS[font], width * 0.96, 0  # 4 % reserve: PowerPoint wraps a little earlier than Chrome
    for para in str(text).split("\n"):
        line, n = "", n + 1
        for word in para.split(" "):
            trial = f"{line} {word}" if line else word
            if f.width(trial, size, spacing) <= limit or not line:
                line = trial
            else:
                n, line = n + 1, word
            while f.width(line, size, spacing) > limit and len(line) > 1:  # a word longer than the box
                n, line = n + 1, line[len(line) // 2:]
    return n


# --- layout model: every slide is a list of px primitives, rendered to pptx or HTML --------------

class Slide:
    def __init__(self, bg, notes=""):
        self.bg, self.notes, self.els = bg, notes, []

    def rect(self, x, y, w, h, fill, r=0, line=None):
        self.els.append(dict(k="rect", x=x, y=y, w=w, h=h, fill=fill, r=r, line=line))

    def image(self, x, y, w, h, path, focus=None, r=0):
        """focus = None stretches to the frame (icons, logo); a focus point crops a photo to fill it."""
        self.els.append(dict(k="image", x=x, y=y, w=w, h=h, path=pathlib.Path(path), focus=focus, r=r))

    def text(self, x, y, w, h, text, size, lh, font="reg", color=C["ink"], anchor="top", align="left", spacing=0):
        self.els.append(dict(k="text", x=x, y=y, w=w, h=h, text=str(text), size=size, lh=lh, font=font,
                             color=color, anchor=anchor, align=align, spacing=spacing))


class Build:
    def __init__(self, deck, base=pathlib.Path(".")):
        self.deck, self.base, self.errors, self.warnings, self.n = deck, base, [], [], 0

    def err(self, msg):
        self.errors.append(f"slide {self.n}: {msg}")

    def fit(self, text, size, width, max_lines, what, font="reg", spacing=0):
        n = count_lines(text, size, width, font, spacing)
        if n > max_lines:
            self.err(f"{what} needs {n} lines, max. {max_lines}: shorten it ({text[:50]!r}…)")
        return min(n, max_lines)

    def trial(self, measure):
        """Runs a layout only to measure it: returns (height, fits the texts) and forgets its errors."""
        mark = len(self.errors)
        h = measure()
        ok = len(self.errors) == mark
        del self.errors[mark:]
        return h, ok

    def count(self, items, lo, hi, what):
        if not lo <= len(items) <= hi:
            self.err(f"{len(items)} {what}, allowed {lo}–{hi}: split the slide or merge items")
        return items[:hi]

    def bottom(self, y):
        if y > CONTENT_BOTTOM:
            self.err(f"content is {y - CONTENT_BOTTOM:.0f} px too tall even in the tightest layout: shorten item texts "
                     "to one line, drop an item or the note, or split the slide")

    # shared pieces

    def photo_el(self, s, x, y, w, h, src, focus=None, r=0):
        """Photo cropped to the frame around its focus point; warns when it is too small to stay sharp."""
        path = next((p for p in (self.base / src, pathlib.Path(src)) if p.is_file()), None)
        if path is None:
            self.err(f"photo not found: {src} (paths are relative to the outline file)")
            return
        f = FOCUS.get(focus or "center") if not isinstance(focus, list) else tuple(focus)
        if f is None or len(f) != 2 or not all(0 <= v <= 1 for v in f):
            self.err(f"focus {focus!r}: use {', '.join(FOCUS)} or [x, y] from 0 to 1 (0, 0 = top left)")
            f = FOCUS["center"]
        from pptx.parts.image import Image
        iw, ih = Image.from_file(str(path)).size
        if min(iw / w, ih / h) < SHARP - .05:
            self.warnings.append(f"slide {self.n}: photo {path.name} is {iw} × {ih} px, this frame needs "
                                 f"{round(w * SHARP)} × {round(h * SHARP)} px to stay sharp; use a larger photo")
        if path.stat().st_size > 3e6 and not shutil.which("sips"):
            try:
                import PIL  # noqa: F401
            except ImportError:
                self.warnings.append(f"slide {self.n}: photo {path.name} is {path.stat().st_size / 1e6:.0f} MB and cannot be "
                                     "scaled down here (pip install pillow); the deck will be heavy")
        s.image(x, y, w, h, shrink(path, w, h), f, r)

    def tile(self, s, x, y, size, icon, variant="red"):
        if icon not in ICONS:
            self.err(f"unknown icon {icon!r}; available: {', '.join(sorted(ICONS))}")
            icon = "check"
        fill, color = {"red": (C["soft"], "red"), "ink": (C["surface"], "ink"), "dark": (C["ink"], "white")}[variant]
        s.rect(x, y, size, size, fill, r=size / 4)
        s.image(x + size / 4, y + size / 4, size / 2, size / 2, SKILL / "icons" / "png" / f"{icon}-{color}.png")

    def item(self, s, x, y, w, it, tile=64, variant="red", ts=(28, 36), ps=(23, 32), gap=24, text_lines=3):
        """Icon tile + title + text; returns the height. With s=None it only measures."""
        tx, tw = x + tile + gap, w - tile - gap
        n = self.fit(it["title"], ts[0], tw, 2, "item title", "med")
        s and s.text(tx, y + 2, tw, n * ts[1], it["title"], *ts, "med")
        h = 2 + n * ts[1]
        if it.get("text"):
            n = self.fit(it["text"], ps[0], tw, text_lines, "item text")
            s and s.text(tx, y + h + 4, tw, n * ps[1], it["text"], *ps, color=C["muted"])
            h += 4 + n * ps[1]
        s and self.tile(s, x, y, tile, it.get("icon", "check"), variant)
        return max(h, tile)

    def stack(self, s, items, x, w, style, gap):
        y = CONTENT_TOP
        for it in items:
            y += self.item(s, x, y, w, it, *style) + gap
        return y - gap

    def grid(self, s, items, cols, style, gap):
        """Items in 2 or 3 columns; the tallest item of a row sets its height so rows line up."""
        col_gap = 80 if cols == 2 else 40
        w = (1120 - col_gap * (cols - 1)) / cols
        y = CONTENT_TOP
        for r in range(0, len(items), cols):
            row = items[r:r + cols]
            h = max(self.item(None, 0, 0, w, it, *style) for it in row)
            for i, it in enumerate(row):
                s and self.item(s, 80 + i * (w + col_gap), y, w, it, *style)
            y += h + gap
        return y - gap

    def layout(self, s, options, limit):
        """Draws the first option whose texts fit and that ends above limit; else the last one, which reports why."""
        for draw in options:
            end, ok = self.trial(lambda: draw(None))
            if ok and end <= limit:
                return draw(s)
        return draw(s)

    def note_h(self, d):
        return 24 + self.trial(lambda: self.fit(d["note"], 22, 1072, 2, "note"))[0] * 30 + 32 if d.get("note") else 0

    def content(self, d, gray_label=False, head_w=960, x=80):
        s = Slide(C["white"], d.get("notes", ""))
        s.text(x, 64, 1200 - x, 24, d.get("label", "").upper(), 18, 24, "med", C["muted"] if gray_label else C["red"], spacing=.1)
        n = self.fit(d["title"], 52, head_w, 2, "headline", "med", -.02)
        s.text(x, 100, head_w, 120, d["title"], 52, 60, "med", spacing=-.02)  # always two lines tall
        return s

    def footer(self, s, left=0, right=1280):
        s.rect(left, 656, right - left, 1, C["border"])
        path, ratio = logo("color")
        s.image(left + 80, 674, 28 * ratio, 28, path)  # 28 px: at 20 px the logo reads as a gray dash
        s.text(left + 80 + 28 * ratio + 16, 676, right - left - 300, 24, self.deck.get("footer", ""), 18, 24, color=C["muted"])
        s.text(right - 180, 676, 100, 24, self.n, 18, 24, color=C["muted"], align="right")

    def note(self, s, d, y):
        if not d.get("note"):
            self.bottom(y)
            return
        n = self.fit(d["note"], 22, 1072, 2, "note")
        y += 24
        s.rect(80, y, 1120, n * 30 + 32, C["surface"], r=12)
        s.text(104, y + 16, 1072, n * 30, d["note"], 22, 30, color=C["muted"])
        self.bottom(y + n * 30 + 32)

    # slide types

    def cover(self, d):
        s = Slide(C["red"], d.get("notes", ""))
        path, ratio = logo("white")
        s.image(80, 64, 40 * ratio, 40, path)
        w = 560 if d.get("photo") else 760
        if d.get("photo"):
            self.photo_el(s, 720, 0, 560, 720, d["photo"], d.get("focus"))
        s.rect(80, 576, w, 1, on_red(.3))
        x = 80
        for part in (d.get("author", ""), d.get("date", "")):
            if part:
                s.text(x, 600, 400, 32, part, 22, 32, color=on_red(.85))
                x += FONTS["reg"].width(part, 22) + 32
        if not d.get("date"):
            self.err("cover has no date")
        y = 528
        if d.get("subtitle"):
            n = self.fit(d["subtitle"], 28, w, 3, "subtitle")
            y -= n * 38
            s.text(80, y, w, n * 38, d["subtitle"], 28, 38, color=on_red(.9))
            y -= 24
        n = self.fit(d["title"], 72, w + 40, 3, "cover title", "med", -.03)
        y -= n * 78
        s.text(80, y, w + 40, n * 78, d["title"], 72, 78, "med", C["white"], spacing=-.03)
        if d.get("label"):
            y -= 44
            s.text(80, y, w, 24, d["label"].upper(), 18, 24, "med", on_red(.8), spacing=.1)
        if y < 128:
            self.err("cover text collides with the logo: shorten the title or subtitle")
        return s

    def list(self, d):
        s = self.content(d)
        intro = d.get("intro")
        items = self.count(d["items"], 1, 4 if intro else 6, "items next to an intro" if intro else "items")
        end, limit = CONTENT_TOP, CONTENT_BOTTOM - self.note_h(d)
        tight = ((REGULAR, 24), (REGULAR, 16), (DENSE, 16))
        if intro:
            n = self.fit(intro, 28, 400, 8, "intro")
            s.text(80, CONTENT_TOP, 400, n * 40, intro, 28, 40)
            end = CONTENT_TOP + n * 40
            options = [lambda t, st=st, g=g: self.stack(t, items, 560, 640, st, g) for st, g in tight]
        else:  # one column while it fits, then 2 columns; 3 only as a last resort (narrow, titles wrap unevenly)
            options = [lambda t, st=st, g=g: self.stack(t, items, 80, 760, st, g) for st, g in tight if len(items) <= 4]
            options += [lambda t, c=c, st=st, g=g: self.grid(t, items, c, st, g) for c in (2, 3) for st, g in tight]
        self.note(s, d, max(end, self.layout(s, options, limit)))
        self.footer(s)
        return s

    def columns(self, d):
        s = self.content(d)
        cols = self.count(d["columns"], 2, 2, "columns")
        for i, c in enumerate(cols):
            x, w = 80 + i * 600, 520
            tx = x
            if c.get("prefix"):
                s.text(x, CONTENT_TOP, 40, 34, c["prefix"], 26, 34, "med", C["g400"])
                tx = x + 40
            self.fit(c["title"], 26, w - (tx - x), 1, "column title", "med")
            s.text(tx, CONTENT_TOP, w - (tx - x), 34, c["title"], 26, 34, "med")
            s.rect(x, CONTENT_TOP + 44, w, 1, C["border"])
        # both columns share the row positions: the taller item of a row sets its height
        lists = [self.count(c["items"], 1, 3, "items in a column") for c in cols]
        def draw(t, style, gap):
            rows = [max(self.item(None, 0, 0, 520, l[r], *style) for l in lists if r < len(l))
                    for r in range(max(map(len, lists)))]
            for i, items in enumerate(lists):
                y = CONTENT_TOP + 68
                for r, it in enumerate(items):
                    t and self.item(t, 80 + i * 600, y, 520, it, *style)
                    y += rows[r] + gap
            return CONTENT_TOP + 68 + sum(rows) + gap * (len(rows) - 1)
        options = [lambda t, st=st, g=g: draw(t, st, g) for st, g in ((COMPACT, 20), (COMPACT, 12), (DENSE, 12))]
        self.note(s, d, self.layout(s, options, CONTENT_BOTTOM - self.note_h(d)))
        self.footer(s)
        return s

    def process(self, d):
        s = self.content(d)
        steps = self.count(d["steps"], 3, 5, "steps")
        n = len(steps)
        cw = (1120 - 24 * (n - 1)) / n
        title_lines = max(self.fit(st["title"], 26, cw, 2, "step title", "med") for st in steps)  # one row for all texts
        end = CONTENT_TOP
        for i, st in enumerate(steps):
            x = 80 + i * (cw + 24)
            if i < n - 1:
                s.rect(x + 76, CONTENT_TOP + 31, cw + 24 - 88, 2, C["border"])
            self.tile(s, x, CONTENT_TOP, 64, st.get("icon", "check"), "dark" if st.get("repeat") else "red")
            s.text(x, 352, cw, 24, f"{i + 1:02d}", 18, 24, "med", C["g400"])
            s.text(x, 380, cw, title_lines * 34, st["title"], 26, 34, "med")
            y = 380 + title_lines * 34 + 6
            if st.get("text"):
                k = self.fit(st["text"], 22, cw, 5, "step text")
                s.text(x, y, cw, k * 30, st["text"], 22, 30, color=C["muted"])
                y += k * 30
            if st.get("note"):
                k = self.fit(st["note"], 18, cw, 2, "step note")
                s.text(x, y + 12, cw, k * 24, st["note"], 18, 24, color=C["muted"])
                y += 12 + k * 24
            end = max(end, y)
        self.note(s, d, end)
        self.footer(s)
        return s

    def statement(self, d):
        s = Slide(C["red"], d.get("notes", ""))
        if d.get("number"):
            s.text(80, 64, 1120, 160, d["number"], 160, 160, "med", on_red(.35), spacing=-.04)
        y = 656
        if d.get("text"):
            n = self.fit(d["text"], 22, 720, 2, "statement text")
            y -= n * 32
            s.text(80, y, 720, n * 32, d["text"], 22, 32, color=on_red(.85))
            y -= 24
        n = self.fit(d["title"], 64, 760, 3, "statement title", "med", -.02)
        y -= n * 72
        s.text(80, y, 760, n * 72, d["title"], 64, 72, "med", C["white"], spacing=-.02)
        return s

    def number(self, d):
        s = self.content(d, gray_label=True)
        self.fit(d["value"], 168, 520, 1, "number", "med", -.04)
        s.text(80, CONTENT_TOP, 560, 168, d["value"], 168, 168, "med", C["red"], spacing=-.04)
        y = CONTENT_TOP + 168 + 16
        if d.get("caption"):
            n = self.fit(d["caption"], 28, 520, 3, "number caption", "med")
            s.text(80, y, 520, n * 36, d["caption"], 28, 36, "med")
            y += n * 36
        if d.get("source"):
            n = self.fit(d["source"], 22, 520, 3, "number source")
            s.text(80, y + 12, 520, n * 30, d["source"], 22, 30, color=C["muted"])
            y += 12 + n * 30
        end = y
        facts = self.count(d.get("facts", []), 0, 3, "facts")
        if facts:
            box = len(s.els)
            y = CONTENT_TOP + 8
            for i, f in enumerate(facts):
                y += 16 + self.item(s, 712, y + 16, 456, f, 56, "ink", *COMPACT[2:]) + 16
                if i < len(facts) - 1:
                    s.rect(712, y, 456, 1, C["border"])
            y += 8
            s.els.insert(box, dict(k="rect", x=680, y=CONTENT_TOP, w=520, h=y - CONTENT_TOP, fill=C["white"], r=16, line=C["border"]))
            end = max(end, y)
        self.bottom(end)
        self.footer(s)
        return s

    def photo(self, d):
        left = d.get("side", "right") == "left"
        x = 720 if left else 80
        s = self.content(d, head_w=480, x=x)
        self.photo_el(s, 0 if left else 640, 0, 640, 720, d["photo"], d.get("focus"))
        if d.get("text"):
            n = self.fit(d["text"], 28, 480, 8, "photo slide text")
            s.text(x, CONTENT_TOP, 480, n * 40, d["text"], 28, 40)
            self.bottom(CONTENT_TOP + n * 40)
        self.footer(s, *((640, 1280) if left else (0, 640)))
        return s

    def photo_full(self, d):
        """The photo is the message: it fills the slide, the claim sits on a dark gradient at the bottom."""
        s = Slide(C["ink"], d.get("notes", ""))
        self.photo_el(s, 0, 0, 1280, 720, d["photo"], d.get("focus"))
        s.image(0, 160, 1280, 560, scrim())
        path, ratio = logo("white")
        s.image(1200 - 32 * ratio, 624, 32 * ratio, 32, path)
        y = 656
        if d.get("text"):
            n = self.fit(d["text"], 24, 720, 2, "photo text")
            y -= n * 34
            s.text(80, y, 720, n * 34, d["text"], 24, 34, color="e8e8e8")
            y -= 16
        n = self.fit(d["title"], 64, 760, 2, "photo title", "med", -.02)
        y -= n * 72
        s.text(80, y, 760, n * 72, d["title"], 64, 72, "med", C["white"], spacing=-.02)
        if d.get("label"):
            s.text(80, y - 44, 760, 24, d["label"].upper(), 18, 24, "med", "e0e0e0", spacing=.1)
        return s

    def gallery(self, d):
        """2–3 photos of one height in a row, each with an optional caption; rounded like other frames."""
        s = self.content(d)
        items = self.count(d["photos"], 2, 3, "photos")
        w = (1120 - 24 * (len(items) - 1)) / len(items)
        lines = max([self.fit(it["caption"], 22, w, 2, "photo caption") for it in items if it.get("caption")] or [0])
        h = CONTENT_BOTTOM - CONTENT_TOP - (12 + lines * 30 if lines else 0)
        for i, it in enumerate(items):
            x = 80 + i * (w + 24)
            self.photo_el(s, x, CONTENT_TOP, w, h, it["photo"], it.get("focus"), r=16)
            if it.get("caption"):
                s.text(x, CONTENT_TOP + h + 12, w, lines * 30, it["caption"], 22, 30, color=C["muted"])
        self.footer(s)
        return s

    def chart(self, d):
        """Horizontal bars: the highlighted item red, the rest gray; the red bar is the hero, so the label is gray."""
        s = self.content(d, gray_label=True)
        bars = self.count(d["bars"], 2, 6, "bars")
        if sum(bool(b.get("highlight")) for b in bars) != 1:
            self.warnings.append(f"slide {self.n}: highlight exactly one bar, the one the headline talks about")
        bottom = CONTENT_BOTTOM
        if d.get("source"):
            n = self.fit(d["source"], 18, 1120, 2, "chart source")
            bottom -= n * 24
            s.text(80, bottom, 1120, n * 24, d["source"], 18, 24, color=C["muted"])
            bottom -= 24
        top = max(b["value"] for b in bars) or 1
        step = min(64, (bottom - CONTENT_TOP - 30) / max(len(bars) - 1, 1))
        unit = d.get("unit", "")
        for i, b in enumerate(bars):
            y = CONTENT_TOP + i * step
            self.fit(b["label"], 22, 300, 1, "bar label")
            s.text(80, y, 300, 30, b["label"], 22, 30, "med" if b.get("highlight") else "reg")
            w = max(4, b["value"] / top * 640)
            s.rect(400, y + 1, w, 28, C["red"] if b.get("highlight") else C["g400"], r=4)
            value = b.get("text") or f"{b['value']:,}".replace(",", " ") + unit
            s.text(400 + w + 16, y, 180, 30, value, 22, 30, "med", C["red"] if b.get("highlight") else C["ink"])
        self.footer(s)
        return s

    def quote(self, d):
        """A client's words on a light gray slide; the quote takes the headline's place, the name sits at the bottom."""
        s = Slide(C["surface"], d.get("notes", ""))
        s.text(80, 64, 1120, 24, d.get("label", "").upper(), 18, 24, "med", C["red"], spacing=.1)
        q = QUOTES.get(self.deck.get("lang", "cs-CZ").split("-")[0], "“”")
        text = q[0] + d["text"] + q[1]
        n = self.fit(text, 44, 960, 5, "quote", "med", -.01)
        s.text(80, 100, 960, max(n, 2) * 56, text, 44, 56, "med", spacing=-.01)
        s.rect(80, 528, 64, 2, C["red"])
        s.text(80, 552, 1120, 34, d["name"], 26, 34, "med")
        if d.get("role"):
            self.fit(d["role"], 22, 1120, 1, "quote role")
            s.text(80, 590, 1120, 30, d["role"], 22, 30, color=C["muted"])
        self.footer(s)
        return s

    def closing(self, d):
        s = Slide(C["red"], d.get("notes", ""))
        path, ratio = logo("white")
        s.image(80, 64, 40 * ratio, 40, path)
        s.rect(80, 560, 1120, 1, on_red(.3))
        x = 80
        if d.get("name"):
            s.text(x, 588, 400, 34, d["name"], 26, 34, "med", C["white"])
            x += FONTS["med"].width(d["name"], 26) + 48
        for icon, key in (("mail", "email"), ("phone", "phone")):
            if d.get(key):
                s.image(x, 591, 28, 28, SKILL / "icons" / "png" / f"{icon}-white.png")
                s.text(x + 40, 588, 400, 34, d[key], 24, 34, color=C["white"])
                x += 40 + FONTS["reg"].width(d[key], 24) + 48
        if x > 1200:
            self.err("contact line is too long for one row")
        y = 512
        if d.get("text"):
            n = self.fit(d["text"], 24, 720, 2, "closing text")
            y -= n * 34
            s.text(80, y, 720, n * 34, d["text"], 24, 34, color=on_red(.9))
            y -= 16
        n = self.fit(d["title"], 72, 760, 3, "closing title", "med", -.03)
        y -= n * 78
        s.text(80, y, 760, n * 78, d["title"], 72, 78, "med", C["white"], spacing=-.03)
        if y < 128:
            self.err("closing text collides with the logo: shorten it")
        return s

    def run(self):
        lang = self.deck.get("lang", "cs-CZ")
        fix = lambda v: typeset(v, lang) if isinstance(v, str) else [fix(i) for i in v] if isinstance(v, list) \
            else {k: fix(i) if k not in ("type", "icon", "photo", "focus", "side", "value", "highlight") else i for k, i in v.items()} if isinstance(v, dict) else v
        self.deck["slides"] = fix(self.deck["slides"])
        types = ("cover", "list", "columns", "process", "statement", "number", "chart", "quote", "photo", "photo-full",
                 "gallery", "closing")
        slides = []
        for self.n, d in enumerate(self.deck["slides"], 1):
            t = d.get("type")
            if t not in types:
                self.err(f"unknown type {t!r}; use one of {', '.join(types)}")
                continue
            try:
                slides.append(getattr(self, t.replace("-", "_"))(d))
            except KeyError as e:
                self.err(f"{t} slide is missing the field {e}")
        self.rhythm([d.get("type") for d in self.deck["slides"]])
        return slides

    def rhythm(self, types):
        red = {"cover", "statement", "closing", "photo-full"}  # heroes: never two in a row
        if types[:1] != ["cover"]:
            self.warnings.append("the deck should open with a red cover")
        if types[-1:] != ["closing"]:
            self.warnings.append("the deck should end with a red closing")
        for i in range(1, len(types)):
            if types[i] in red and types[i - 1] in red:
                self.errors.append(f"slides {i}–{i + 1}: two hero slides in a row (red or full photo); put a white slide between")
        for i in range(2, len(types)):
            if types[i] == types[i - 1] == types[i - 2] and types[i] not in red:
                self.warnings.append(f"slides {i - 1}–{i + 1}: three {types[i]} slides in a row; alternate the types")
        if types.count("statement") > max(1, len(types) // 8):
            self.warnings.append("too many red statement slides: one per 8 slides")
        if types.count("photo-full") > 2:
            self.warnings.append("more than two full-slide photos: they stop standing out; use photo or gallery")


# --- renderers ------------------------------------------------------------------------------------

def to_pptx(slides, out, lang):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Pt

    E = lambda v: Emu(round(v * EMU))
    rgb = RGBColor.from_string
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(9144000), Emu(5143500)
    theme = prs.slide_master.part.part_related_by(RT.THEME)  # theme fonts -> Inter, so nothing falls back to Calibri
    theme._blob = re.sub(rb'(<a:majorFont>\s*<a:latin typeface=")[^"]*', rb"\1Inter Medium", theme.blob)
    theme._blob = re.sub(rb'(<a:minorFont>\s*<a:latin typeface=")[^"]*', rb"\1Inter", theme._blob)
    for sl in slides:
        p = prs.slides.add_slide(prs.slide_layouts[6])
        p.background.fill.solid()
        p.background.fill.fore_color.rgb = rgb(sl.bg)
        for el in sl.els:
            x, y, w, h = E(el["x"]), E(el["y"]), E(el["w"]), E(el["h"])
            if el["k"] == "rect":
                shp = p.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if el["r"] else MSO_SHAPE.RECTANGLE, x, y, w, h)
                if el["r"]:
                    shp.adjustments[0] = el["r"] / min(el["w"], el["h"])
                shp.fill.solid()
                shp.fill.fore_color.rgb = rgb(el["fill"])
                if el["line"]:
                    shp.line.color.rgb = rgb(el["line"])
                    shp.line.width = E(1)
                else:
                    shp.line.fill.background()
                style = shp._element.find(qn("p:style"))  # drop the theme style (shadow, accent text)
                if style is not None:
                    shp._element.remove(style)
            elif el["k"] == "image":
                pic = p.shapes.add_picture(str(el["path"]), x, y, w, h)
                if el["focus"]:  # crop to the frame's ratio around the focus point, never stretch
                    pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = crop(*pic.image.size, el["w"], el["h"], el["focus"])
                if el["r"]:
                    geom = pic._element.spPr.find(qn("a:prstGeom"))
                    geom.set("prst", "roundRect")
                    av = geom.find(qn("a:avLst"))
                    av.append(av.makeelement(qn("a:gd"), {"name": "adj", "fmla": f"val {round(el['r'] / min(el['w'], el['h']) * 100000)}"}))
            else:
                tf = p.shapes.add_textbox(x, y, w, h).text_frame
                tf.word_wrap, tf.auto_size = True, MSO_AUTO_SIZE.NONE
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "bottom": MSO_ANCHOR.BOTTOM}[el["anchor"]]
                for i, para in enumerate(el["text"].split("\n")):
                    par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    par.alignment = PP_ALIGN.RIGHT if el["align"] == "right" else PP_ALIGN.LEFT
                    par.line_spacing = Pt(el["lh"] * PT)
                    run = par.add_run()
                    run.text = para
                    f = run.font
                    f.name, f.size, f.bold = FAMILY[el["font"]], Pt(el["size"] * PT), False
                    f.color.rgb = rgb(el["color"])
                    rpr = run._r.get_or_add_rPr()
                    rpr.set("lang", lang)
                    if el["spacing"]:
                        rpr.set("spc", str(round(el["size"] * PT * el["spacing"] * 100)))
        if sl.notes:
            p.notes_slide.notes_text_frame.text = sl.notes
    prs.save(out)


def to_html(slides, out, lang):
    css = []
    for key, fname in (("reg", "Inter-Regular.ttf"), ("med", "Inter-Medium.ttf")):
        b64 = base64.b64encode((SKILL / "fonts" / fname).read_bytes()).decode()
        css.append(f'@font-face{{font-family:"{FAMILY[key]}";src:url(data:font/ttf;base64,{b64})}}')
    css.append("body{margin:0;background:#d9d9d9;padding:24px;display:grid;gap:24px;justify-content:center}"
               ".s{position:relative;width:1280px;height:720px;overflow:hidden}.s>*{position:absolute;margin:0}"
               ".t{display:flex;flex-direction:column;white-space:pre-wrap}")
    html = [f'<!doctype html><html lang="{lang}"><meta charset="utf-8"><title>Deck preview</title><style>{"".join(css)}</style>']
    for sl in slides:
        html.append(f'<div class="s" style="background:#{sl.bg}">')
        for el in sl.els:
            pos = f'left:{el["x"]}px;top:{el["y"]}px;width:{el["w"]}px;height:{el["h"]}px'
            if el["k"] == "rect":
                line = f';box-shadow:inset 0 0 0 1px #{el["line"]}' if el["line"] else ""
                html.append(f'<div style="{pos};background:#{el["fill"]};border-radius:{el["r"]}px{line}"></div>')
            elif el["k"] == "image":
                fit = f'cover;object-position:{el["focus"][0] * 100}% {el["focus"][1] * 100}%' if el["focus"] else "fill"
                html.append(f'<img src="{el["path"].resolve().as_uri()}" style="{pos};object-fit:{fit};border-radius:{el["r"]}px">')
            else:
                text = el["text"].replace("&", "&amp;").replace("<", "&lt;")
                html.append(f'<div class="t" style="{pos};font:{el["size"]}px/{el["lh"]}px \'{FAMILY[el["font"]]}\';'
                            f'color:#{el["color"]};letter-spacing:{el["spacing"]}em;text-align:{el["align"]};'
                            f'justify-content:{"flex-end" if el["anchor"] == "bottom" else "flex-start"}">{text}</div>')
        html.append("</div>")
    pathlib.Path(out).write_text("\n".join(html))


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("outline", help="deck outline, JSON")
    ap.add_argument("out", help="output .pptx")
    ap.add_argument("--preview", help="also write an HTML preview of the same layout")
    a = ap.parse_args()
    deck = json.loads(pathlib.Path(a.outline).read_text())
    b = Build(deck, pathlib.Path(a.outline).resolve().parent)
    slides = b.run()
    for w in b.warnings:
        print("warning:", w, file=sys.stderr)
    b.errors = list(dict.fromkeys(b.errors))  # measured and drawn items report the same problem twice
    if b.errors:
        print("\n".join(["Deck not built, fix the outline:"] + [f"  - {e}" for e in b.errors]), file=sys.stderr)
        sys.exit(1)
    lang = deck.get("lang", "cs-CZ")
    to_pptx(slides, a.out, lang)
    if a.preview:
        to_html(slides, a.preview, lang.split("-")[0])
    print(f"{a.out}: {len(slides)} slides" + (f", preview {a.preview}" if a.preview else ""))
    print("Inter must be installed where the .pptx is opened (the skill's fonts/), or PowerPoint swaps the font and "
          "line breaks move. For clients send the PDF: python3 scripts/check_deck.py deck.pptx --render DIR")


if __name__ == "__main__":
    main()
